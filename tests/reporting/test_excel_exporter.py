import unittest
import pandas as pd
from pathlib import Path
from openpyxl import load_workbook
from quantbullet.reporting.excel_exporter import ExcelExporter
from quantbullet.reporting.columns import ColumnFormat, ColumnMeta, ColumnSchema
from quantbullet.reporting.converters import convert_pdf_schema_to_excel_schema
from quantbullet.reporting.pdf_text_report import PdfTextReport, PdfColumnFormat, PdfColumnMeta
from tests.artifacts import artifact_dir

class TestExcelExporter(unittest.TestCase):
    def setUp(self):
        self.cache_dir = artifact_dir(self, "reporting/excel_exporter")

        df = pd.DataFrame({
            "A": [1, 2, 3],
            "B": [4.5678, 5.6789, 6.7890],
            "C": ["foo", "bar", "baz"],
            "D": ["2024-01-01", "2024-01-02", "2024-01-03"],
            "E": ["abc", "def", "ghi"]
        })

        df["D"] = pd.to_datetime( df["D"] )
        self.df = df
        return df

    def path(self, name):
        return str(Path(self.cache_dir) / name)

    def column_a(self, path, sheet):
        return [cell.value for cell in load_workbook(path)[sheet]["A"]]

    def test_excel_exporter_basic(self):
        df = self.df
        exporter = ExcelExporter( filename=str( Path(self.cache_dir) / "test_exporter.xlsx" ) )
        schema = ColumnSchema( columns=[
            ColumnMeta( name="A", display_name="Column A", format=ColumnFormat( decimals=0 ) ),
            ColumnMeta( name="B", display_name="Column B", format=ColumnFormat( decimals=2, comma=True, transformer=lambda x: x * 1000 ) ),
            ColumnMeta( name="C", display_name="Column C", format=ColumnFormat() ),
            ColumnMeta( name="D", display_name="Date Column", format=ColumnFormat() )
        ] )
        exporter.add_sheet( sheet_name="TestSheet", df = df, schema=schema )
        exporter.save()

    def test_pdf_text_report( self ):
        df = self.df
        report = PdfTextReport( file_path=str( Path(self.cache_dir) / "test_report.pdf" ) )
        schema = [
            PdfColumnMeta(name="A", display_name="Column A", format=PdfColumnFormat(decimals=0)),
            PdfColumnMeta(name="B", display_name="Column B", format=PdfColumnFormat(decimals=2, comma=True, transformer=lambda x: x * 1000)),
            PdfColumnMeta(name="C", display_name="Column C", format=PdfColumnFormat()),
            PdfColumnMeta(name="D", display_name="Date Column", format=PdfColumnFormat())
        ]
        report.add_df_table( df=df, schema=schema )
        report.save()

        excel_schema = convert_pdf_schema_to_excel_schema( schema )
        exporter = ExcelExporter( filename=str( Path(self.cache_dir) / "test_report_converted.xlsx" ) )
        exporter.add_sheet( sheet_name="TestSheet", df = df, schema=excel_schema )
        exporter.save()

    def test_schema_with_default_formats(self):
        path = self.path("defaults.xlsx")
        schema = ColumnSchema([ColumnMeta(name="A"), ColumnMeta(name="B", format=ColumnFormat(decimals=None))])
        ExcelExporter(path).add_sheet("s", self.df, schema=schema).save()
        ws = load_workbook(path)["s"]
        self.assertEqual([ws["A1"].value, ws["B1"].value], ["A", "B"])
        self.assertEqual(ws["A2"].number_format, "0.00")
        self.assertEqual(ws["B2"].number_format, "General")

    def test_schema_column_missing_from_df(self):
        schema = ColumnSchema([ColumnMeta(name="Z")])
        with self.assertRaisesRegex(ValueError, "'Z'"):
            ExcelExporter(self.path("missing.xlsx")).add_sheet("s", self.df, schema=schema)

    def test_width_is_estimated_from_transformed_values(self):
        path = self.path("width.xlsx")
        schema = ColumnSchema([ColumnMeta(name="v", format=ColumnFormat(transformer=lambda x: x * 1e6))])
        ExcelExporter(path).add_sheet("s", pd.DataFrame({"v": [1.0, 2.0]}), schema=schema).save()
        self.assertEqual(load_workbook(path)["s"].column_dimensions["A"].width, len("2000000.00") + 2)

    def test_tables_stack_on_one_sheet(self):
        path = self.path("stack.xlsx")
        df = pd.DataFrame({"a": [1, 2]})
        ExcelExporter(path).add_sheet("s", df).add_sheet("s", df * 10).add_sheet("s", df * 100, gap_rows=0).save()
        self.assertEqual(self.column_a(path, "s"), ["a", 1, 2, None, "a", 10, 20, "a", 100, 200])

    def test_tables_on_one_sheet_of_existing_file(self):
        path = self.path("existing.xlsx")
        df = pd.DataFrame({"a": [1, 2]})
        ExcelExporter(path).add_sheet("keep", df).add_sheet("s", df * 5).save()
        ExcelExporter(path).add_sheet("s", df).add_sheet("s", df * 10, start_row=5).save()
        self.assertEqual(load_workbook(path).sheetnames, ["keep", "s"])
        self.assertEqual(self.column_a(path, "s"), ["a", 1, 2, None, "a", 10, 20])
        self.assertEqual(self.column_a(path, "keep"), ["a", 1, 2])

    def test_overwrite_drops_other_sheets(self):
        path = self.path("overwrite.xlsx")
        df = pd.DataFrame({"a": [1]})
        ExcelExporter(path).add_sheet("old", df).save()
        ExcelExporter(path).set_overwrite(True).add_sheet("new", df).save()
        self.assertEqual(load_workbook(path).sheetnames, ["new"])

    def test_failure_raises_original_error_and_keeps_file_and_queue(self):
        path = self.path("failure.xlsx")
        df = pd.DataFrame({"a": [1]})
        ExcelExporter(path).add_sheet("keep", df).save()
        before = Path(path).read_bytes()

        exporter = ExcelExporter(path).set_overwrite(True).add_sheet("bad/name", df)
        with self.assertRaisesRegex(ValueError, "Invalid character"):
            exporter.save()
        self.assertEqual(Path(path).read_bytes(), before)
        self.assertEqual(len(exporter._blocks), 1)

    def test_freeze_panes(self):
        path = self.path("freeze.xlsx")
        df = pd.DataFrame({"a": [1]})
        ExcelExporter(path).add_sheet("s", df).add_sheet("t", df).set_freeze_panes("s", "B2").save()
        workbook = load_workbook(path)
        self.assertEqual(workbook["s"].freeze_panes, "B2")
        self.assertIsNone(workbook["t"].freeze_panes)

    def test_freeze_panes_validation(self):
        with self.assertRaisesRegex(ValueError, "Invalid cell"):
            ExcelExporter(self.path("bad.xlsx")).set_freeze_panes("s", "2B")
        exporter = ExcelExporter(self.path("bad.xlsx")).add_sheet("s", self.df).set_freeze_panes("other", "A2")
        with self.assertRaisesRegex(ValueError, "other"):
            exporter.save()

    def test_font(self):
        path = self.path("font.xlsx")
        df = pd.DataFrame({"value": [1.0]})
        ExcelExporter(path).set_font(name="Arial", size=22, color="1F4E79").add_sheet("s", df).save()
        ws = load_workbook(path)["s"]
        self.assertEqual((ws["A1"].font.name, ws["A1"].font.sz, ws["A1"].font.b), ("Arial", 22, True))
        self.assertEqual((ws["A2"].font.name, ws["A2"].font.b, ws["A2"].font.color.rgb), ("Arial", False, "001F4E79"))
        self.assertEqual(ws.column_dimensions["A"].width, (len("value") + 2) * 2)

    def test_save_without_tables(self):
        with self.assertRaisesRegex(ValueError, "No tables"):
            ExcelExporter(self.path("empty.xlsx")).save()
