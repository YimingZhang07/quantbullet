"""ExcelExporter: one minimal yet complete workbook that uses each capability once, and how re-saving updates it.

``test_minimal_complete_workbook`` writes one workbook, in one font, with two sheets:

- ``Loans``: a schema that picks, renames, orders and formats columns (decimals, thousands separators, percent,
  negatives in parentheses, a transformer, a date format, Excel's General format, a fixed width, a formula and
  color scales, a column alignment), with a wrapped header, and its header row and ID column frozen.
- ``Summary``: several tables on one sheet, stacked below each other, a Total row appended without a header, and
  a table without a schema placed beside them at an explicit position.

``test_update_one_sheet`` saves a workbook twice: the second save replaces only the sheet it names.

``setUpClass`` builds one fake loan tape that both examples start from. The numbers are synthetic; only the layout
matters. Run an example with the run button next to its method in the VS Code Test panel (``examples`` must be
listed in ``python.testing.pytestArgs``), or from the repository root:

    .venv/Scripts/python.exe -m pytest examples/reporting/test_excel_exporter_examples.py -k minimal_complete

The workbooks are written to OUTPUT_DIR, which git ignores; open them in Excel to see the formats.
"""
import unittest
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from quantbullet.reporting import ColumnFormat, ColumnMeta, ColumnSchema, ExcelExporter

# Where the examples write their workbooks. Change it to write somewhere else.
OUTPUT_DIR = Path(__file__).resolve().parent / "output" / "excel_exporter"
RATINGS = ["AAA", "AA", "A", "BBB"]


def fake_loans(n_rows: int, seed: int) -> pd.DataFrame:
    """A loan tape: id, origination date and vintage, rating, balance ($), price (per 100), rate and spread
    (decimals, e.g. 0.0525 and 0.0150), P&L ($, about half negative), default probability and a model score."""
    rng = np.random.default_rng(seed)
    orig_date = pd.Timestamp("2021-01-01") + pd.to_timedelta(rng.integers(0, 3 * 365, n_rows), unit="D")
    return pd.DataFrame({
        "loan_id": [f"L{i:03d}" for i in range(1, n_rows + 1)],
        "orig_date": orig_date,
        "vintage": orig_date.year,
        "rating": rng.choice(RATINGS, n_rows),
        "balance": rng.uniform(50_000, 750_000, n_rows).round(-2),
        "price": rng.uniform(90, 105, n_rows),
        "rate": rng.uniform(0.03, 0.08, n_rows),
        "spread": rng.uniform(0.01, 0.04, n_rows),
        "pnl": rng.normal(0, 5_000, n_rows),
        "default_prob": rng.uniform(0, 0.2, n_rows),
        "score": rng.normal(0, 1, n_rows),
    })


def output_path(name: str) -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    return OUTPUT_DIR / name


class TestExcelExporterExamples(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.loans = fake_loans(n_rows=40, seed=0)   # both examples start from this loan tape

    def test_minimal_complete_workbook(self):
        """Start here: every column format and table placement, in one workbook."""
        path = output_path("minimal_complete_workbook.xlsx")

        # ---- Loans: one table, every column format.
        loans = self.loans.assign(market_value=lambda d: d["balance"] * d["price"] / 100)
        loans_schema = ColumnSchema([                     # only these columns are written, in this order: vintage is dropped
            ColumnMeta("loan_id", "Loan ID", ColumnFormat(width=8)),          # width: fixed; None fits values and header
            ColumnMeta("orig_date", "Origination Date",
                       ColumnFormat(date_format="mmm yyyy")),                  # an Excel format code, not strftime
            ColumnMeta("rating", format=ColumnFormat(align="center")),         # display_name defaults to the name;
                                                                               # align: any value type, overrides set_alignment
            ColumnMeta("balance", "Balance", ColumnFormat(decimals=0, comma=True)),   # comma: thousands separators
            ColumnMeta("price", "Price", ColumnFormat(decimals=3)),
            ColumnMeta("market_value", "Market Value",
                       ColumnFormat(decimals=0, comma=True,
                                    formula_template="=D{row}*E{row}/100")),   # a live formula per row replaces the values
            ColumnMeta("rate", "Rate", ColumnFormat(decimals=2, percent=True)),        # Excel shows 0.0525 as 5.25%
            ColumnMeta("spread", "Spread (bp)",
                       ColumnFormat(decimals=0, transformer=lambda x: x * 1e4)),       # the values written are transformed
            ColumnMeta("pnl", "P&L", ColumnFormat(decimals=0, comma=True,
                                                  parens_for_negative=True,           # -1,234 shows as (1,234)
                                                  color_scale=True)),                 # red (low) to green (high)
            ColumnMeta("default_prob", "Default Probability",
                       ColumnFormat(decimals=1, percent=True, color_scale=True,
                                    higher_is_better=False)),                         # reversed: high is red
            ColumnMeta("score", "Model Score", ColumnFormat(decimals=None)),          # None: Excel's General format
        ])

        # ---- Summary: several tables on one sheet.
        measures = [
            ColumnMeta("count", "Loans", ColumnFormat(decimals=0)),
            ColumnMeta("balance", "Balance", ColumnFormat(decimals=0, comma=True)),
            ColumnMeta("rate", "Average Rate", ColumnFormat(decimals=2, percent=True)),
        ]
        aggregations = dict(count=("loan_id", "size"), balance=("balance", "sum"), rate=("rate", "mean"))
        by_rating = self.loans.groupby("rating", as_index=False).agg(**aggregations)
        total = pd.DataFrame({"rating": ["Total"], "count": [len(self.loans)],
                              "balance": [self.loans["balance"].sum()], "rate": [self.loans["rate"].mean()]})
        by_vintage = self.loans.groupby("vintage", as_index=False).agg(**aggregations)
        balance_mm = (self.loans.pivot_table(index="rating", columns="vintage", values="balance", aggfunc="sum") / 1e6
                      ).rename(columns=str).reset_index()                    # balance ($mm) by rating x vintage
        rating_schema = ColumnSchema([ColumnMeta("rating", "Rating"), *measures])
        vintage_schema = ColumnSchema([ColumnMeta("vintage", "Vintage", ColumnFormat(decimals=0)), *measures])

        (
            ExcelExporter(path)
            .set_overwrite(True)                                 # start a new workbook; False (default) keeps other sheets
            .set_font(name="Arial", size=10)                     # every written cell, on all sheets; None keeps Calibri 11
            .set_alignment(None)                                 # every data cell, e.g. "right"; None: text left, numbers right
            .add_sheet("Loans", loans, schema=loans_schema,
                       wrap_header=True)                         # wrap long headers in a taller header row
            .set_freeze_panes("Loans", "B2")                     # rows above 2 and column A stay put; "A2" freezes the header only
            .add_sheet("Summary", by_rating, schema=rating_schema,
                       start_row=None)                           # None: the first table on a sheet starts at row 1 ...
            .add_sheet("Summary", total, schema=rating_schema,
                       include_header=False,                     # no header row: the Total row continues the table
                       gap_rows=0)                               # ... and later ones go below, gap_rows blank rows apart
            .add_sheet("Summary", by_vintage, schema=vintage_schema, gap_rows=2)
            .add_sheet("Summary", balance_mm,
                       schema=None,                              # no schema: every column, numbers with 2 decimals
                       start_row=1, start_col=7)                 # an explicit position: beside the first table
            .save()                                              # writes the file only if every table succeeded
        )
        print(f"written: {path}")

        workbook = load_workbook(path)
        self.assertEqual(workbook.sheetnames, ["Loans", "Summary"])
        loans_sheet, summary = workbook["Loans"], workbook["Summary"]
        self.assertEqual(loans_sheet.freeze_panes, "B2")
        self.assertEqual(loans_sheet["C2"].alignment.horizontal, "center")
        self.assertEqual(summary["A2"].font.name, "Arial")
        self.assertEqual(loans_sheet["F2"].value, "=D2*E2/100")
        self.assertEqual(loans_sheet["I2"].number_format, "#,##0;(#,##0);#,##0")
        self.assertEqual(summary.cell(row=len(by_rating) + 2, column=1).value, "Total")
        self.assertEqual(summary.cell(row=len(by_rating) + 5, column=1).value, "Vintage")
        self.assertEqual(summary["G1"].value, "rating")

    def test_update_one_sheet(self):
        """Saving to an existing workbook replaces the sheets it names and keeps the others."""
        path = output_path("update_one_sheet.xlsx")
        columns = ["loan_id", "rating", "balance"]
        (
            ExcelExporter(path)
            .set_overwrite(True)
            .add_sheet("Loans", self.loans[columns])
            .add_sheet("Ratings", pd.DataFrame({"rating": RATINGS}))
            .save()
        )
        # A later run that refreshes Loans only: overwrite=False (the default) keeps Ratings.
        ExcelExporter(path).set_overwrite(False).add_sheet("Loans", self.loans[columns].head(5)).save()
        print(f"written: {path}")

        workbook = load_workbook(path)
        self.assertEqual(workbook.sheetnames, ["Loans", "Ratings"])
        self.assertEqual(workbook["Loans"].max_row, 1 + 5)


if __name__ == "__main__":
    unittest.main()
